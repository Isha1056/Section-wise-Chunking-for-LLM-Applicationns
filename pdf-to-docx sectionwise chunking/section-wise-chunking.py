import win32com.client
import os
import json
from datetime import datetime 
word = win32com.client.Dispatch("Word.Application")
word.visible = 1
filename = "currency_internationisation"
pdfdoc = os.path.join(os.getcwd(),f'PDF/{filename}.pdf')
todocx = os.path.join(os.getcwd(),f'DOCX/{filename}.docx')
wb1 = word.Documents.Open(pdfdoc)
wb1.SaveAs(todocx, FileFormat=16)  # file format for docx
wb1.Close()
word.Quit()

import docx
import docx.text.run
# d = docx.Document(todocx)
d = docx.Document(todocx)
def write(cur, res): # cur for depth
    x = res
    for i in range(len(cur)):
        if cur[i][0] not in x:
            if cur[i][1] == 100:
                x[cur[i][0]] = None
                break
            else:
                x[cur[i][0]] = {}
        x = x.get(cur[i][0])
res = {}
cur = []
for p in d.paragraphs:
    # print(p.style.name, " , ", p.text)
    p_text = p.text.strip()
    f = 0                   
    for run in p.runs:
        if run.bold or run.italic or run.underline:
            f = 1
    if not p_text:
        continue
    if 'Heading' not in p.style.name and f == 0: # no formatting and no heading
        level = 100
        if len(cur) and cur[-1][1] == 100:
            cur[-1][0] = cur[-1][0] + " " + p_text
        else:
            cur.append([p_text, level])
    else:
        if 'Heading' not in p.style.name: # formatting but no heading
            level = 11
            x = ''
            for run in p.runs:
                if run.bold or run.italic or run.underline:
                    if run.text.strip():
                        x += run.text.strip() + ' '
            if p_text.startswith(x.strip()):
                if len(cur) and level <= cur[-1][1]:
                    write(cur, res)
                while len(cur) and level <= cur[-1][1]:
                    cur.pop()
                cur.append([x, level])
            level = 100
            x = ''
            for run in p.runs:
                if not run.bold and not run.italic and not run.underline:
                    if run.text.strip():
                        x += run.text.strip() + ' '
            if len(cur) and cur[-1][1] == 100:
                cur[-1][0] = cur[-1][0] + " " + x
            else:
                cur.append([x, level])
        else:  # Headings
            level = int(p.style.name.split()[-1])
            if len(cur) and level <= cur[-1][1]:
                write(cur, res)
            while len(cur) and level <= cur[-1][1]:
                cur.pop()
            cur.append([p_text, level])
if len(cur):
    x = res
    write(cur, res)
print(json.dumps(res))
l = []

def dfs(x, s):
    if x == None:
        l.append(["|".join(s)])
    else:
        for k,v in x.items():
            dfs(v, s+[k])
dfs(res,[])

import csv
out_file = open(os.getcwd()+f"/CSV/{filename}.csv", "w", newline='', encoding='utf-8')
writer = csv.writer(out_file)
writer.writerows(l)
out_file.close()